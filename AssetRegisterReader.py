from openpyxl import load_workbook
from typing import List, Dict, Tuple
import json, os


class AssetRegister:
    def __init__(self, filename: str = "NA Asset Register 2023.xlsx"):
        data_dir = os.getenv("DATA_DIR", ".")
        filename = filename if os.path.isabs(filename) else os.path.join(data_dir, filename)
        # Check if the "NA Asset Register 2023.xlsx" file exists, if not then use any other .xlsx file in the same folder
        if os.path.exists(filename):
            self.register = load_workbook(filename, read_only=True)
        else:
            # Find any other .xlsx file in the same folder
            folder = os.getenv("DATA_DIR", os.path.dirname(filename) or ".")
            xlsx_files = [f for f in os.listdir(folder) if f.endswith(".xlsx")]
            
            if xlsx_files:
                fallback_file = os.path.join(folder, xlsx_files[0])
                print(f"Default file not found. Using fallback file: {fallback_file}")
                self.register = load_workbook(fallback_file, read_only=True)
            else:
                raise FileNotFoundError("No .xlsx file found in the folder.")

    def get_assets(self):
        """
        Get all assets' data.
        Returns a list of dictionaries. Each dictionary will be in the following format:
        
        ``{
            "asset_number": "",
            "sap_number": "",
            "name": "",
            "present_location": "",
            "condition": ""
        }``
        """
        sheets = self.register.sheetnames
        results: List[Dict[str, str]] = [] # Will store the details of all the assets

        for sheet_name in sheets:
            current_sheet = self.register[sheet_name]

            # Rows doesn't start from 1, so we need to find the row where the titles are
            if current_sheet["A12"].value and current_sheet["A12"].value.strip() == "#":
                titles_row = 12
            elif current_sheet["A13"].value and current_sheet["A13"].value.strip() == "#":
                titles_row = 13
            elif current_sheet["A14"].value and current_sheet["A14"].value.strip() == "#":
                titles_row = 14
            elif current_sheet["A15"].value and current_sheet["A15"].value.strip() == "#":
                titles_row = 15
            elif current_sheet["A16"].value and current_sheet["A16"].value.strip() == "#":
                titles_row = 16

            titles: List[Tuple[str, str]] = []
            for row in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"]:
                title: str | None = current_sheet[f"{row}{titles_row}"].value
                if title:
                    titles.append((title.lower().strip(), row))

            start_row = titles_row + 1

            for row in range(start_row, 100):
                data = {
                    "asset_number": "",
                    "sap_number": "",
                    "name": "",
                    "present_location": "",
                    "condition": ""
                }
                for each_title, column  in titles:
                    if each_title == "asset number":
                        asset_num_cell: str = current_sheet[f"{column}{row}"].value
                        if asset_num_cell:
                            data["asset_number"] = asset_num_cell.strip()

                    elif each_title in ["sap asset number", "sap asset no."]:
                        sap_num_cell: str | int = current_sheet[f"{column}{row}"].value
                        if sap_num_cell:
                            data["sap_number"] = str(sap_num_cell).strip()

                    elif "name" in each_title:
                        name_cell: str = current_sheet[f"{column}{row}"].value
                        if name_cell:
                            data["name"] = name_cell.strip()

                    elif "present" in each_title:
                        present_loc_cell: str = current_sheet[f"{column}{row}"].value
                        if present_loc_cell:
                            data["present_location"] = present_loc_cell.strip()

                    elif "condition" in each_title:
                        condition_cell: str = current_sheet[f"{column}{row}"].value
                        if condition_cell:
                            data["condition"] = condition_cell.strip()

                if data != {
                    "asset_number": "",
                    "sap_number": "",
                    "name": "",
                    "present_location": "",
                    "condition": ""
                }:
                    results.append(data)
                else:
                    break
        
        return results

    def export_to_json(self):
        """Exports the assets' information into a JSON file called `assets.json`"""
        data = self.get_assets()
        assets_path = os.path.join(os.getenv("DATA_DIR", "."), "assets.json")
        with open(assets_path, "w", encoding="utf-8") as json_file:
            json.dump(data, json_file, indent=4)
        
        return
    
    def read_assets_json_file(self):
        """Reads the data from the exported `assets.json` file"""
        try:
            assets_path = os.path.join(os.getenv("DATA_DIR", "."), "assets.json")
            with open(assets_path, "r", encoding="utf-8") as json_file:
                data: List[Dict[str, str]] = json.load(json_file)
            return data
        except FileNotFoundError:
            return []
    
if __name__ == "__main__":
    ar = AssetRegister()
    ar.read_assets_json_file()

            